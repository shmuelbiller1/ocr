"""Self-test support.

The Windows build is only considered good if the packaged executable
actually starts.  ``hebrew_workbench.selftest`` runs a small number of
staged checks inside the real application and writes a JSON report, so
the CI job can fail loudly instead of assuming that a produced ``.exe``
file means a working application.

The self-test is written to survive a broken installation: every stage is
guarded, the report file is written even when a stage crashes, and the
process always ends with a meaningful exit code.
"""

from __future__ import annotations

import json
import platform
import sys
import traceback
from dataclasses import dataclass, field
from pathlib import Path

__all__ = ["CheckResult", "run_checks", "main"]


@dataclass
class CheckResult:
    name: str
    ok: bool
    detail: str = ""

    def to_dict(self) -> dict:
        return {"name": self.name, "ok": self.ok, "detail": self.detail}


@dataclass
class _Report:
    results: list[CheckResult] = field(default_factory=list)

    def add(self, name: str, ok: bool, detail: str = "") -> None:
        self.results.append(CheckResult(name=name, ok=ok, detail=detail))

    @property
    def ok(self) -> bool:
        return all(result.ok for result in self.results)


def _write_line(text: str) -> None:
    """Print without ever raising (a windowed build has no console)."""
    try:
        print(text)
    except Exception:
        pass


def _check_imports(report: _Report) -> None:
    modules = [
        "hebrew_workbench",
        "hebrew_workbench.engine.base",
        "hebrew_workbench.engine.registry",
        "hebrew_workbench.project",
    ]
    failures = []
    for name in modules:
        try:
            __import__(name)
        except Exception as exc:  # pragma: no cover - defensive
            failures.append(f"{name}: {exc}")
    report.add(
        "core-imports",
        not failures,
        "; ".join(failures) if failures else "all core modules import cleanly",
    )


def _check_engine_contract(report: _Report) -> None:
    import tempfile

    from .engine.base import EngineUnavailable, OCRResult
    from .engine.registry import EngineRegistry

    problems = []
    registry = EngineRegistry()
    if registry.engine.available():
        problems.append("the default engine must not report itself as available")
    try:
        registry.transcribe(Path("does-not-exist.png"))
    except FileNotFoundError:
        pass
    except Exception as exc:
        problems.append(f"missing files must raise FileNotFoundError, got {exc!r}")

    # A real file on disk is needed: inside a packaged build ``__file__``
    # points into the archive and does not exist as a path.
    with tempfile.TemporaryDirectory() as folder:
        sample = Path(folder) / "sample.png"
        sample.write_bytes(b"not really a png")
        try:
            registry.transcribe(sample)
        except EngineUnavailable:
            pass
        except Exception as exc:
            problems.append(
                f"an unavailable engine must raise EngineUnavailable, got {exc!r}"
            )
        else:
            problems.append("an unavailable engine must raise EngineUnavailable")

        class _Fake:
            name = "self-test adapter"

            def available(self) -> bool:
                return True

            def recognize(self, file_path: Path) -> OCRResult:
                return OCRResult("שלום", self.name)

        try:
            result = EngineRegistry(_Fake()).transcribe(sample)
        except Exception as exc:
            problems.append(f"a working adapter must produce a result, got {exc!r}")
        else:
            if result.text != "שלום" or result.engine_name != "self-test adapter":
                problems.append("a working adapter must return its own result untouched")
    report.add(
        "engine-contract",
        not problems,
        "; ".join(problems) if problems else "the engine contract fails closed",
    )


def _check_gui(report: _Report) -> None:
    try:
        from PySide6.QtWidgets import QApplication

        from .ui.main_window import MainWindow
    except Exception as exc:  # pragma: no cover - defensive
        report.add(
            "gui-startup",
            False,
            f"the user interface could not be imported: {exc}\n{traceback.format_exc()}",
        )
        return

    app = QApplication.instance() or QApplication([])
    try:
        window = MainWindow()
        window.show()
        app.processEvents()
        visible = window.isVisible()
        title = window.windowTitle()
        window.close()
        app.processEvents()
    except Exception as exc:  # pragma: no cover - defensive
        report.add(
            "gui-startup",
            False,
            f"the main window could not be created: {exc}\n{traceback.format_exc()}",
        )
        return
    report.add(
        "gui-startup",
        bool(visible) and bool(title),
        f"main window created with title {title!r}",
    )


def _check_preview(report: _Report) -> None:
    """Make sure the document preview really works in this build.

    This is the check that catches a packaging mistake: when PDFium or
    Pillow is missing from the executable, the preview fails while
    everything else still looks fine.
    """
    import base64
    import tempfile

    try:
        from .preview.loader import DocumentPreview, PreviewError
    except Exception as exc:  # pragma: no cover - defensive
        report.add("preview", False, f"the preview module could not be imported: {exc}")
        return

    problems = []
    # A 1x1 PNG, so that the check does not depend on Pillow being able to
    # write a file.
    tiny_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmM"
        "IQAAAABJRU5ErkJggg=="
    )
    with tempfile.TemporaryDirectory() as folder:
        sample = Path(folder) / "tiny.png"
        sample.write_bytes(tiny_png)
        try:
            with DocumentPreview(sample) as preview:
                page = preview.render(1)
                if page.image.width < 1:
                    problems.append("the rendered page is empty")
        except PreviewError as exc:
            problems.append(f"an image could not be previewed: {exc}")
        except Exception as exc:
            problems.append(f"the image preview crashed: {exc!r}")

    from .resources import samples_directory

    samples = samples_directory()
    if samples is None:
        report.add(
            "preview",
            not problems,
            "; ".join(problems)
            if problems
            else "images preview correctly (the sample folder is not bundled here)",
        )
        return

    pdf = samples / "two_pages.pdf"
    if pdf.exists():
        try:
            with DocumentPreview(pdf) as preview:
                if preview.page_count != 2:
                    problems.append(
                        f"the sample PDF should have two pages, it has {preview.page_count}"
                    )
                preview.render(2)
        except Exception as exc:
            problems.append(f"the sample PDF could not be previewed: {exc!r}")
    else:
        problems.append("the sample PDF is missing from the bundle")

    report.add(
        "preview",
        not problems,
        "; ".join(problems) if problems else "images and PDF pages preview correctly",
    )


def _check_plugins(report: _Report) -> None:
    """Make sure the plugin system works in this build.

    The check installs a small plugin into a temporary folder, loads it,
    and takes it away again.  A build in which the plugin system does
    not work is a build in which another engine cannot be added.
    """
    import json
    import tempfile

    try:
        from .plugins import PluginInstaller, PluginRegistry
    except Exception as exc:  # pragma: no cover - defensive
        report.add("plugins", False, f"the plugin system could not be imported: {exc}")
        return

    problems: list[str] = []
    try:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "plugins"
            registry = PluginRegistry(root=root)
            ids = {plugin.id for plugin in registry.plugins}
            from .plugins.builtin import BUILTIN_MANIFESTS

            expected = {manifest.id for manifest in BUILTIN_MANIFESTS}
            missing = expected - ids
            if missing:
                problems.append(f"missing built-in plugins: {sorted(missing)}")
            # Missing optional weights/worker may legitimately produce None;
            # an import or configuration error is never a healthy plugin.
            for builtin in registry.plugins:
                if builtin.error:
                    problems.append(f"{builtin.id}: {builtin.error}")

            source = Path(folder) / "example"
            source.mkdir()
            (source / "plugin.json").write_text(
                json.dumps(
                    {
                        "id": "selftest-engine",
                        "name": "Selftest engine",
                        "version": "0.0.1",
                        "kind": "ocr",
                        "entry": "selftest_engine:SelftestEngine",
                    }
                ),
                encoding="utf-8",
            )
            (source / "selftest_engine.py").write_text(
                "class SelftestEngine:\n"
                "    name = 'Selftest engine'\n"
                "    def available(self):\n"
                "        return True\n",
                encoding="utf-8",
            )
            installer = PluginInstaller(root)
            installer.install_folder(source)
            registry = PluginRegistry(root=root)
            plugin = registry.get("selftest-engine")
            if plugin is None or plugin.object is None:
                problems.append("an installed plugin was not loaded")
            installer.remove("selftest-engine")
    except Exception as exc:  # pragma: no cover - defensive
        problems.append(f"the plugin system failed: {exc!r}")

    report.add(
        "plugins",
        not problems,
        "; ".join(problems) if problems else "plugins are installed and loaded",
    )


def _check_document_output(report: _Report) -> None:
    """Exercise real project/TXT/DOCX/PDF output, not OCR or a physical printer.

    The synthetic text is explicitly a test fixture. Use the main window's
    printer factory so obsolete Qt APIs cannot pass a launch-only check.
    """
    import tempfile

    import docx
    from PySide6.QtPrintSupport import QPrinter
    from PySide6.QtWidgets import QApplication

    from .exporting import build_print_document, print_document, write_docx, write_txt
    from .preview.loader import DocumentPreview
    from .project import Project
    from .resources import samples_directory
    from .ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    try:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            samples = samples_directory()
            if samples is None:
                raise RuntimeError("bundled sample pages are missing")
            sample = samples / "scan_page_1.png"
            before = sample.read_bytes()
            raw = "בדיקת פלט 12345\nשורה שנייה English"
            edited = raw + "\nעריכה לבדיקה"
            project = Project()
            entry = project.add(sample)
            entry.original_transcription = raw
            entry.transcription = edited
            restored = Project.load(project.save(root / "check.hwproj"))
            if (restored.documents[0].transcription != edited or
                    restored.documents[0].original_transcription != raw):
                raise RuntimeError("project did not preserve original and edited text")
            txt = write_txt(root / "check.txt", edited)
            if txt.read_text(encoding="utf-8-sig") != edited:
                raise RuntimeError("TXT output changed the text")
            word = docx.Document(str(write_docx(root / "check.docx", edited)))
            if "\n".join(p.text for p in word.paragraphs) != edited:
                raise RuntimeError("DOCX output changed the text")
            if not all('w:bidi' in p._p.xml for p in word.paragraphs):
                raise RuntimeError("DOCX paragraphs are not marked right-to-left")
            pdf = root / "check.pdf"
            printer = window._printer()
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(str(pdf))
            print_document(build_print_document(edited), printer)
            if not pdf.read_bytes().startswith(b"%PDF-"):
                raise RuntimeError("PDF output is missing or invalid")
            with DocumentPreview(pdf) as preview:
                if preview.page_count < 1:
                    raise RuntimeError("PDF output has no pages")
                preview.render(1)
            if sample.read_bytes() != before:
                raise RuntimeError("the original scan was modified")
    finally:
        window.close()
        window.deleteLater()
        app.processEvents()
    report.add("document-output", True,
               "synthetic fixture: project/raw text, TXT, RTL DOCX and rendered PDF; "
               "physical printer output is not tested")


def run_checks() -> list[CheckResult]:
    """Run every staged check and return the individual results."""
    report = _Report()
    report.add(
        "environment",
        True,
        f"python {platform.python_version()} on {platform.platform()}",
    )
    for name, check in (
        ("core-imports", _check_imports),
        ("engine-contract", _check_engine_contract),
        ("preview", _check_preview),
        ("plugins", _check_plugins),
        ("gui-startup", _check_gui),
        ("document-output", _check_document_output),
    ):
        try:
            check(report)
        except Exception as exc:  # pragma: no cover - defensive
            report.add(name, False, f"the check crashed: {exc!r}")
    return report.results


def _write_report(path: str | Path, payload: dict) -> bool:
    try:
        target = Path(path)
        if target.parent and str(target.parent) not in ("", "."):
            target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return True
    except Exception:
        return False


def main(report_path: str | Path | None = None) -> int:
    """Run the self-test and optionally write a JSON report.

    Returns the process exit code: ``0`` when every check passed.
    """
    results: list[CheckResult] = []
    crash = ""
    try:
        results = run_checks()
    except Exception as exc:  # pragma: no cover - defensive
        crash = f"{exc!r}\n{traceback.format_exc()}"
        results = [CheckResult("selftest", False, f"the self-test crashed: {crash}")]

    payload = {
        "ok": all(result.ok for result in results) and not crash,
        "python": sys.version,
        "results": [result.to_dict() for result in results],
    }
    if crash:
        payload["crash"] = crash

    if report_path:
        _write_report(report_path, payload)

    _write_line(json.dumps(payload, ensure_ascii=False, indent=2))
    for result in results:
        status = "PASS" if result.ok else "FAIL"
        _write_line(f"[{status}] {result.name}: {result.detail}")
    return 0 if payload["ok"] else 1


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1 else None))
