"""Writing the transcription out.

Three output paths are supported, all of them local:

* plain UTF-8 text (``.txt``),
* Word documents (``.docx``) with right-to-left paragraphs,
* printing, either to a Windows printer or to a PDF file.

None of them changes the text: what the user sees in the editor is what
gets written out.
"""

from __future__ import annotations

from pathlib import Path

__all__ = [
    "ExportError",
    "write_txt",
    "write_docx",
    "build_print_document",
    "print_document",
]


class ExportError(RuntimeError):
    """The file could not be written."""


def write_txt(path: str | Path, text: str) -> Path:
    """Write plain text with a UTF-8 BOM so Windows editors show Hebrew."""
    target = Path(path)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8-sig")
    except OSError as exc:
        raise ExportError(f"לא ניתן לשמור את הקובץ: {exc}") from exc
    return target


def _set_paragraph_rtl(paragraph) -> None:
    """Mark a Word paragraph as right-to-left."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    properties = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    properties.append(bidi)
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), "right")
    properties.append(justification)
    for run in paragraph.runs:
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        run._r.get_or_add_rPr().append(rtl)


def write_docx(path: str | Path, text: str) -> Path:
    """Write a Word document with right-to-left paragraphs."""
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExportError(
            "חבילת ייצוא Word (python-docx) אינה מותקנת, ולכן לא ניתן לייצא קובץ DOCX."
        ) from exc

    target = Path(path)
    try:
        document = docx.Document()
        style = document.styles["Normal"]
        style.font.name = "David"
        style.font.size = Pt(14)
        lines = text.splitlines() or [""]
        for line in lines:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(line)
            run.font.size = Pt(14)
            _set_paragraph_rtl(paragraph)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(target))
    except Exception as exc:
        raise ExportError(f"לא ניתן לשמור את קובץ ה-Word: {exc}") from exc
    return target


def build_print_document(text: str, font_family: str = "David", font_size: int = 14):
    """Build a print-ready Qt document with right-to-left paragraphs."""
    from PySide6.QtCore import Qt
    from PySide6.QtGui import (
        QFont,
        QTextBlockFormat,
        QTextCharFormat,
        QTextCursor,
        QTextDocument,
    )

    document = QTextDocument()
    document.setDefaultFont(QFont(font_family, font_size))

    cursor = QTextCursor(document)
    block_format = QTextBlockFormat()
    block_format.setLayoutDirection(Qt.RightToLeft)
    block_format.setAlignment(Qt.AlignRight)
    char_format = QTextCharFormat()
    char_format.setFont(QFont(font_family, font_size))

    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        cursor.setBlockFormat(block_format)
        cursor.setCharFormat(char_format)
        cursor.insertText(line)
        if index < len(lines) - 1:
            cursor.insertBlock()
    return document


def print_document(document, printer) -> None:
    """Send a prepared document to a printer."""
    try:
        document.print_(printer)
    except Exception as exc:  # pragma: no cover - depends on the printer
        raise ExportError(f"ההדפסה נכשלה: {exc}") from exc
