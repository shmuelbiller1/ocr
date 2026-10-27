"""Tests for the export and print helpers."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from hebrew_workbench.exporting import (  # noqa: E402
    ExportError,
    build_print_document,
    write_docx,
    write_txt,
)

SAMPLE = "זהוי טקסט לבדיקה\nשורה שנייה"


class ExportTests(unittest.TestCase):
    def test_txt_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "out.txt"
            write_txt(target, SAMPLE)
            self.assertTrue(target.exists())
            self.assertEqual(target.read_text(encoding="utf-8-sig"), SAMPLE)

    def test_txt_is_written_with_a_byte_order_mark(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "out.txt"
            write_txt(target, SAMPLE)
            self.assertTrue(target.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_docx_contains_the_text_and_is_right_to_left(self) -> None:
        import docx

        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "out.docx"
            write_docx(target, SAMPLE)
            self.assertTrue(target.exists())
            document = docx.Document(str(target))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("זהוי טקסט לבדיקה", text)
            xml = document.paragraphs[0]._p.xml
            self.assertIn("w:bidi", xml)

    def test_print_document_holds_the_text(self) -> None:
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])
        document = build_print_document(SAMPLE)
        self.assertIn("זהוי טקסט לבדיקה", document.toPlainText())

    def test_empty_text_is_allowed(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "empty.txt"
            write_txt(target, "")
            self.assertTrue(target.exists())

    def test_unwritable_path_raises(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            # A directory cannot be overwritten by a file.
            target = Path(folder) / "folder"
            target.mkdir()
            with self.assertRaises(ExportError):
                write_txt(target, "text")


if __name__ == "__main__":
    unittest.main()
